"""Leitura da planilha de despesas exportada do ERP.

Funcao pura: entra um arquivo (caminho ou stream), sai uma lista de linhas
validas e uma lista de erros posicionados. Nao importa Flask nem toca no banco.

O export tem o cabecalho na linha 2, intercala linhas de subtotal entre os
lancamentos, exporta valores negativos e datas como serial do Excel.
"""

from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

COLUNAS = {
    "ANO_BAIXA",
    "MES_BAIXA",
    "DIA_BAIXA",
    "DATAEMISSAO",
    "REFERENCIA",
    "FORNECEDOR",
    "CR_REDUZIDO",
    "NATUREZA",
    "HISTORICO",
    "DOCUMENTO",
    "VALOR ORIGINAL",
    "VALOR BAIXADO",
}

COLUNAS_ESSENCIAIS = {"REFERENCIA", "FORNECEDOR", "VALOR BAIXADO"}
LINHAS_ATE_O_CABECALHO = 10

CENTAVO = Decimal("0.01")
_ESPACOS = re.compile(r"\s+")

# O serial 60 e 29/02/1900, dia que nunca existiu: o Excel herdou o bug do
# Lotus 1-2-3. Por isso a epoca muda de lado no meio.
_EPOCA_ANTES_DO_BUG = date(1899, 12, 31)
_EPOCA_DEPOIS_DO_BUG = date(1899, 12, 30)


@dataclass(frozen=True)
class LinhaPlanilha:
    linha: int
    referencia: int
    data_baixa: date
    data_emissao: date | None
    fornecedor: str
    centro_custo: str
    natureza: str
    historico: str
    documento: str
    valor_original: Decimal
    valor_baixado: Decimal

    @property
    def divergente(self) -> bool:
        return self.valor_original != self.valor_baixado


@dataclass(frozen=True)
class ErroLinha:
    linha: int
    campo: str
    mensagem: str


@dataclass
class Resultado:
    linhas: list[LinhaPlanilha] = field(default_factory=list)
    erros: list[ErroLinha] = field(default_factory=list)
    ignoradas: int = 0

    @property
    def ok(self) -> bool:
        return not self.erros


def normalizar_texto(valor: object) -> str:
    """Colapsa espacos e quebras de linha, e apara as pontas."""
    if valor is None:
        return ""
    return _ESPACOS.sub(" ", str(valor)).strip()


def serial_para_data(serial: float | int) -> date:
    dias = int(serial)
    if dias <= 0:
        raise ValueError(f"serial de data inválido: {serial!r}")
    if dias == 60:
        raise ValueError("serial 60 é 29/02/1900, data que nunca existiu")
    if dias < 60:
        return _EPOCA_ANTES_DO_BUG + timedelta(days=dias)
    return _EPOCA_DEPOIS_DO_BUG + timedelta(days=dias)


def eh_linha_de_total(celulas: list, mapa: dict | None = None) -> bool:
    """Rotulo terminando em ' Total' e ultima coisa preenchida antes dos valores.
    So o sufixo derrubaria um fornecedor chamado "... Total"."""
    limite = _inicio_dos_valores(celulas, mapa)
    rotulos = celulas[:limite]

    while rotulos and _celula_vazia(rotulos[-1]):
        rotulos.pop()

    if not rotulos:
        return False

    # Sem coluna vazia a direita nao da para distinguir de um DOCUMENTO que
    # termine em " Total"; preferimos manter o lancamento.
    if len(rotulos) >= limite:
        return False

    ultimo = rotulos[-1]
    return isinstance(ultimo, str) and ultimo.rstrip().endswith(" Total")


def _celula_vazia(valor) -> bool:
    return valor is None or (isinstance(valor, str) and not valor.strip())


def _inicio_dos_valores(celulas: list, mapa: dict | None) -> int:
    if mapa:
        posicoes = [mapa[nome] for nome in ("VALOR ORIGINAL", "VALOR BAIXADO") if nome in mapa]
        if posicoes:
            return min(posicoes)
    return len(celulas)


def ler(origem) -> Resultado:
    """Le a planilha. `origem` pode ser um caminho ou um stream."""
    livro = _abrir(origem)
    try:
        linhas = list(livro.worksheets[0].iter_rows(values_only=True))
    finally:
        livro.close()

    indice_cabecalho, mapa = _localizar_cabecalho(linhas)
    resultado = Resultado()

    for deslocamento, celulas in enumerate(linhas[indice_cabecalho + 1 :]):
        numero = indice_cabecalho + deslocamento + 2
        celulas = list(celulas)

        if eh_linha_de_total(celulas, mapa):
            resultado.ignoradas += 1
            continue
        if _vazia(celulas):
            continue

        _ler_linha(numero, celulas, mapa, resultado)

    _rejeitar_lancamentos_repetidos(resultado)
    return resultado


def chave_natural(linha: LinhaPlanilha) -> tuple:
    """Fornecedor e natureza sem caixa, como o obter_ou_criar casa."""
    return (
        linha.referencia,
        linha.data_baixa,
        linha.fornecedor.upper(),
        linha.natureza.upper(),
        linha.centro_custo,
        linha.documento,
        linha.historico,
    )


def _rejeitar_lancamentos_repetidos(resultado: Resultado) -> None:
    """Duas linhas iguais na chave natural quebrariam o commit inteiro."""
    vistas: dict[tuple, int] = {}
    repetidas: set[tuple] = set()
    for linha in resultado.linhas:
        chave = chave_natural(linha)
        if chave in vistas:
            repetidas.add(chave)
            resultado.erros.append(
                ErroLinha(
                    linha=linha.linha,
                    campo="REFERENCIA",
                    mensagem=(
                        f"lançamento repetido: referência {linha.referencia} com o mesmo "
                        f"fornecedor, natureza, centro de custo, documento e histórico "
                        f"da linha {vistas[chave]}"
                    ),
                )
            )
        else:
            vistas[chave] = linha.linha

    if repetidas:
        resultado.linhas = [
            linha for linha in resultado.linhas if chave_natural(linha) not in repetidas
        ]


def _abrir(origem):
    try:
        return load_workbook(origem, read_only=True, data_only=True)
    except (zipfile.BadZipFile, InvalidFileException, OSError, ValueError, KeyError) as erro:
        raise ValueError(
            f"o arquivo não parece uma planilha .xlsx válida ({erro.__class__.__name__})"
        ) from erro


def _localizar_cabecalho(linhas: list) -> tuple[int, dict[str, int]]:
    """O cabecalho e procurado, e nao assumido: uma linha a mais no topo nao
    pode virar 127 erros de leitura."""
    for indice, celulas in enumerate(linhas[:LINHAS_ATE_O_CABECALHO]):
        mapa = {}
        for posicao, valor in enumerate(celulas):
            nome = normalizar_texto(valor).upper()
            if nome in COLUNAS:
                mapa[nome] = posicao
        if COLUNAS_ESSENCIAIS.issubset(mapa):
            return indice, mapa

    esperadas = ", ".join(sorted(COLUNAS_ESSENCIAIS))
    raise ValueError(
        f"não encontrei o cabeçalho da planilha nas primeiras "
        f"{LINHAS_ATE_O_CABECALHO} linhas (esperava as colunas {esperadas})"
    )


def _vazia(celulas: list) -> bool:
    return all(valor is None or str(valor).strip() == "" for valor in celulas)


def _ler_linha(numero: int, celulas: list, mapa: dict[str, int], resultado: Resultado) -> None:
    erros: list[ErroLinha] = []

    def bruto(campo: str):
        posicao = mapa.get(campo)
        if posicao is None or posicao >= len(celulas):
            return None
        return celulas[posicao]

    def falhar(campo: str, mensagem: str) -> None:
        erros.append(ErroLinha(linha=numero, campo=campo, mensagem=mensagem))

    referencia = _inteiro(bruto("REFERENCIA"), "REFERENCIA", falhar)
    data_baixa = _data_de_baixa(bruto("ANO_BAIXA"), bruto("MES_BAIXA"), bruto("DIA_BAIXA"), falhar)
    data_emissao = _data_qualquer(bruto("DATAEMISSAO"), "DATAEMISSAO", falhar)
    fornecedor = normalizar_texto(bruto("FORNECEDOR"))
    if not fornecedor:
        falhar("FORNECEDOR", "vazio")
    valor_original = _dinheiro(bruto("VALOR ORIGINAL"), "VALOR ORIGINAL", falhar)
    valor_baixado = _dinheiro(bruto("VALOR BAIXADO"), "VALOR BAIXADO", falhar)

    if erros:
        # Uma linha ruim nao derruba as boas.
        resultado.erros.extend(erros)
        return

    resultado.linhas.append(
        LinhaPlanilha(
            linha=numero,
            referencia=referencia,
            data_baixa=data_baixa,
            data_emissao=data_emissao,
            fornecedor=fornecedor,
            centro_custo=normalizar_texto(bruto("CR_REDUZIDO")),
            natureza=normalizar_texto(bruto("NATUREZA")),
            historico=normalizar_texto(bruto("HISTORICO")),
            documento=normalizar_texto(bruto("DOCUMENTO")),
            valor_original=valor_original,
            valor_baixado=valor_baixado,
        )
    )


def _inteiro(valor, campo: str, falhar) -> int | None:
    if valor is None or str(valor).strip() == "":
        falhar(campo, "vazio")
        return None
    try:
        return int(float(str(valor).strip()))
    except (TypeError, ValueError):
        falhar(campo, f"não é um número inteiro: {valor!r}")
        return None


def _dinheiro(valor, campo: str, falhar) -> Decimal | None:
    """Decimal positivo com dois decimais.

    Positivo porque o ERP exporta tudo negativo e aqui tudo e despesa. Decimal
    porque o arquivo ja traz -2556.0500000000002.
    """
    if valor is None or str(valor).strip() == "":
        falhar(campo, "vazio")
        return None
    try:
        bruto = Decimal(_normalizar_numero(valor))
    except (InvalidOperation, TypeError, ValueError):
        falhar(campo, f"não é um valor numérico: {valor!r}")
        return None
    return bruto.copy_abs().quantize(CENTAVO, rounding=ROUND_HALF_UP)


def _normalizar_numero(valor) -> str:
    """Aceita 1234.56 e o 1.234,56 que se digita no Brasil.

    O export do ERP manda numero de verdade, mas planilha preenchida a mao vem
    com virgula decimal — e recusar isso faria o proprio modelo de importacao
    ser rejeitado.
    """
    # Celula numerica ja vem desambiguada pelo Excel: o ponto e decimal, ponto
    # final. Aplicar heuristica de milhar aqui transformaria 1234.567 em
    # 1234567.
    if isinstance(valor, (int, float, Decimal)) and not isinstance(valor, bool):
        return str(valor)

    texto = str(valor).strip().replace("\xa0", "").replace(" ", "")
    if "," in texto:
        # Com virgula presente, o ponto so pode ser separador de milhar.
        return texto.replace(".", "").replace(",", ".")

    # Texto com so ponto e ambiguo. Em dinheiro digitado a mao, ponto seguido de
    # exatamente 3 digitos e milhar ("1.234"), e mais de um ponto idem.
    partes = texto.lstrip("-+").split(".")
    if len(partes) > 2 or (len(partes) == 2 and len(partes[1]) == 3):
        return texto.replace(".", "")
    return texto


def _data_de_baixa(ano, mes, dia, falhar) -> date | None:
    partes = []
    for valor, campo in ((ano, "ANO_BAIXA"), (mes, "MES_BAIXA"), (dia, "DIA_BAIXA")):
        if valor is None or str(valor).strip() == "":
            falhar("DATA_BAIXA", f"{campo} vazio")
            return None
        try:
            partes.append(int(float(str(valor).strip())))
        except (TypeError, ValueError):
            falhar("DATA_BAIXA", f"{campo} não é número: {valor!r}")
            return None
    try:
        return date(*partes)
    except ValueError as erro:
        falhar("DATA_BAIXA", f"data inválida ({partes[2]:02d}/{partes[1]:02d}/{partes[0]}): {erro}")
        return None


def _data_qualquer(valor, campo: str, falhar) -> date | None:
    """DATAEMISSAO chega como serial, datetime ou texto. Ausente nao e erro."""
    if valor is None or str(valor).strip() == "":
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, int | float):
        try:
            return serial_para_data(valor)
        except ValueError as erro:
            falhar(campo, str(erro))
            return None

    texto = str(valor).strip()
    try:
        return serial_para_data(float(texto))
    except ValueError:
        pass
    for formato in ("%d/%m/%Y", "%Y-%m-%d", "%d/%m/%Y %H:%M:%S"):
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            continue
    falhar(campo, f"não consegui interpretar como data: {valor!r}")
    return None
